"""Integration checks for the migrated Excel recalculation tool."""
from pathlib import Path
import hashlib
import json
import subprocess
import sys
import tempfile
from openpyxl import Workbook, load_workbook

script = Path(__file__).resolve().parent.parent / "apps/zcode-cli/packages/bundled-skills/skills/xlsx/scripts/recalc.py"
with tempfile.TemporaryDirectory(prefix="jzx-recalc-test-") as directory:
    path = Path(directory) / "test.xlsx"
    workbook = Workbook()
    workbook.active.append([80, 90])
    workbook.active["C1"] = "=AVERAGE(A1:B1)"
    workbook.active["C2"] = '=""'
    workbook.save(path)
    result = subprocess.run([sys.executable, str(script), str(path)], capture_output=True, text=True, check=True)
    assert json.loads(result.stdout)["total_formulas"] == 2
    assert load_workbook(path, data_only=True).active["C1"].value == 85
    broken = Path(directory) / "invalid.xlsx"
    broken.write_bytes(b"invalid workbook")
    before = hashlib.sha256(broken.read_bytes()).hexdigest()
    failure = subprocess.run([sys.executable, str(script), str(broken)], capture_output=True, text=True)
    assert failure.returncode != 0
    assert "error" in json.loads(failure.stdout)
    assert hashlib.sha256(broken.read_bytes()).hexdigest() == before
print("Excel integration passed: real cache = 85; failed input preserved")
