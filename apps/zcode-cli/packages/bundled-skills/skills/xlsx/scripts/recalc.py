"""Recalculate XLSX formulas with LibreOffice and verify cached results."""
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from openpyxl import load_workbook
from office.soffice import get_soffice_env


def recalc(filename, timeout=30):
    source = Path(filename).resolve()
    if not source.is_file():
        return {"error": "Excel file does not exist"}
    if source.suffix.lower() != ".xlsx":
        return {"error": "Recalculation requires an .xlsx file"}
    errors = {code: [] for code in ["#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A"]}
    try:
        # 原宏路径即使未保存也返回0；隔离转换并核验缓存，防止假成功覆盖教师文件。
        with tempfile.TemporaryDirectory(prefix="jzx-recalc-", dir=source.parent) as temporary:
            root = Path(temporary)
            output = root / "output"
            output.mkdir()
            profile = (root / "profile").as_uri()
            command = ["soffice", "-env:UserInstallation=" + profile, "--headless",
                       "--convert-to", "xlsx", "--outdir", str(output), str(source)]
            result = subprocess.run(command, env=get_soffice_env(), capture_output=True,
                                    text=True, timeout=timeout, check=False)
            generated = output / source.name
            if result.returncode != 0 or not generated.is_file():
                return {"error": "LibreOffice did not produce a recalculated workbook"}
            formulas = load_workbook(generated, data_only=False)
            values = load_workbook(generated, data_only=True)
            missing = []
            total = 0
            for sheet in formulas:
                for row in sheet:
                    for cell in row:
                        if cell.data_type != "f":
                            continue
                        total += 1
                        cached = values[sheet.title][cell.coordinate]
                        value = cached.value
                        location = f"{sheet.title}!{cell.coordinate}"
                        if value is None and cached.data_type not in ("s", "str"):
                            missing.append(location)
                        elif isinstance(value, str) and value in errors:
                            errors[value].append(location)
            formulas.close()
            values.close()
            if missing:
                return {"error": "Formula caches are missing; original file preserved",
                        "missing_caches": missing[:20], "total_formulas": total}
            summary = {code: {"count": len(locations), "locations": locations[:20]}
                       for code, locations in errors.items() if locations}
            count = sum(len(locations) for locations in errors.values())
            # 输出和源位于同一文件系统；通过核验后一次替换，不写入半成品。
            if count == 0:
                shutil.copymode(source, generated)
                os.replace(generated, source)
            return {"status": "errors_found" if count else "success", "total_errors": count,
                    "error_summary": summary, "total_formulas": total}
    except subprocess.TimeoutExpired:
        return {"error": "LibreOffice recalculation timed out; original file preserved"}
    except Exception as error:
        return {"error": str(error)}


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python recalc.py <excel_file> [timeout_seconds]")
        sys.exit(1)
    response = recalc(sys.argv[1], int(sys.argv[2]) if len(sys.argv) > 2 else 30)
    print(json.dumps(response, indent=2))
    sys.exit(1 if "error" in response or response.get("status") == "errors_found" else 0)
